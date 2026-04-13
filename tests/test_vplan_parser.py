# Copyright (c) 2026 BelTech Systems LLC and contributors
# SPDX-License-Identifier: MIT
"""tests/test_vplan_parser.py — Unit and integration tests for parser/vplan_parser.py.

Phase: v6a
Layer: Tests

Tests VplanParseResult construction from synthetic openpyxl workbooks (unit)
and from the real BALU VPR fixture (integration). Verifies duck-type
compatibility with both ReqParseResult and IntentParseResult.
"""
import os
import tempfile

import openpyxl
import pytest

from parser.vplan_parser import parse_vplan, VplanParseResult

# ── Fixture path ──────────────────────────────────────────────────────────────
_BALU_VPLAN = os.path.join(
    os.path.dirname(__file__), "fixtures", "balu_vplan.xlsx"
)


# ── Synthetic workbook helpers ────────────────────────────────────────────────

def _make_workbook(vpr_rows: list[list], cov_rows: list[list]) -> str:
    """Build a minimal VPR workbook and return its path.

    VPR sheet layout:
      Row 1: group header placeholders (skipped by parser)
      Row 2: column header placeholders (skipped by parser)
      Row 3: example row placeholder (skipped by parser)
      Row 4+: data rows supplied in vpr_rows

    Coverage_Goals sheet layout:
      Row 1: column header (skipped by parser)
      Row 2+: data rows supplied in cov_rows

    Each vpr_row must supply 19+ values corresponding to columns A..S (0-based
    indices 0-18). Missing trailing columns are padded with None.

    Each cov_row must supply 9+ values for columns A..I (0-based 0-8).
    """
    wb = openpyxl.Workbook()

    # ── VPR sheet ──
    ws_vpr = wb.active
    ws_vpr.title = "VPR"
    # Row 1: group headers (skip)
    ws_vpr.append(["Group headers"] + [""] * 33)
    # Row 2: column headers (skip)
    ws_vpr.append(
        ["Req_ID", "Family", "Col_C", "Statement", "Col_E", "Col_F",
         "Verification_Method", "Covered_By", "Col_I", "Col_J",
         "Disposition", "Waiver_Rationale", "Col_M", "Col_N", "Col_O",
         "Col_P", "Col_Q", "Col_R", "RTL_Status"]
    )
    # Row 3: template example (skip)
    ws_vpr.append(["[BLOCK-FAM-NNN]"] + [""] * 18)
    # Data rows
    for row in vpr_rows:
        # Pad to 19 columns
        padded = list(row) + [None] * (19 - len(row))
        ws_vpr.append(padded)

    # ── Coverage_Goals sheet ──
    ws_cov = wb.create_sheet("Coverage_Goals")
    # Row 1: headers (skip)
    ws_cov.append(
        ["ID", "Name", "Description", "Stimulus_Strategy",
         "Boundary_Values", "Linked_Requirements", "Status",
         "Coverage_Type", "Notes"]
    )
    for row in cov_rows:
        padded = list(row) + [None] * (9 - len(row))
        ws_cov.append(padded)

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb.save(tmp.name)
    return tmp.name


def _simple_vpr_row(
    req_id: str,
    family: str = "FAM",
    statement: str = "Stmt",
    verif: str = "simulation",
    covered_by: str = "COV-001",
    disposition: str = "",
    waiver_reason: str = "",
    rtl_status: str = "",
) -> list:
    """Build a 19-element VPR data row."""
    row = [None] * 19
    row[0]  = req_id
    row[1]  = family
    row[3]  = statement
    row[6]  = verif
    row[7]  = covered_by
    row[10] = disposition
    row[11] = waiver_reason
    row[18] = rtl_status
    return row


def _simple_cov_row(
    cov_id: str,
    name: str = "CovName",
    description: str = "Desc",
    stimulus: str = "Stim",
    boundary: str = "",
    linked: str = "",
    status: str = "Open",
    cov_type: str = "Functional",
    notes: str = "",
) -> list:
    """Build a 9-element Coverage_Goals data row."""
    return [cov_id, name, description, stimulus, boundary,
            linked, status, cov_type, notes]


# ── Unit tests ────────────────────────────────────────────────────────────────

def test_parse_vplan_returns_vplan_parse_result() -> None:
    """parse_vplan() returns a VplanParseResult instance."""
    path = _make_workbook([], [])
    try:
        result = parse_vplan(path)
        assert isinstance(result, VplanParseResult)
    finally:
        os.unlink(path)


def test_all_data_rows_in_requirements() -> None:
    """All VPR data rows (including waived) appear in requirements dict."""
    rows = [
        _simple_vpr_row("UART-BR-001"),
        _simple_vpr_row("UART-BR-002", disposition="WAIVED", waiver_reason="OOS"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert "UART-BR-001" in result.requirements
        assert "UART-BR-002" in result.requirements
    finally:
        os.unlink(path)


def test_requirement_entry_fields() -> None:
    """Requirement entries contain expected field keys with correct values."""
    rows = [_simple_vpr_row(
        "UART-BR-001",
        family="BR",
        statement="Counter must reset",
        verif="simulation",
        covered_by="COV-002",
        rtl_status="PASS",
    )]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        entry = result.requirements["UART-BR-001"]
        assert entry["statement"] == "Counter must reset"
        assert entry["verification"] == ["simulation"]
        assert entry["waived"] is False
        assert entry["waiver_reason"] == ""
        assert entry["family"] == "BR"
        assert entry["covered_by"] == "COV-002"
        assert entry["rtl_status"] == "PASS"
    finally:
        os.unlink(path)


def test_verification_method_splits_on_comma() -> None:
    """Comma-separated Verification_Method is split into a list."""
    rows = [_simple_vpr_row("UART-BR-001", verif="simulation, post-silicon, inspection")]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert result.requirements["UART-BR-001"]["verification"] == [
            "simulation", "post-silicon", "inspection"
        ]
    finally:
        os.unlink(path)


def test_waived_row_fields() -> None:
    """Waived row has waived=True and waiver_reason populated."""
    rows = [_simple_vpr_row(
        "UART-PAR-003",
        disposition="WAIVED",
        waiver_reason="Not applicable to this device variant",
    )]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        entry = result.requirements["UART-PAR-003"]
        assert entry["waived"] is True
        assert entry["waiver_reason"] == "Not applicable to this device variant"
    finally:
        os.unlink(path)


def test_waivers_list_contains_waived_ids_only() -> None:
    """waivers list contains waived IDs; non-waived IDs are absent."""
    rows = [
        _simple_vpr_row("UART-BR-001"),
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED", waiver_reason="OOS"),
        _simple_vpr_row("UART-BR-002"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert "UART-PAR-003" in result.waivers
        assert "UART-BR-001" not in result.waivers
        assert "UART-BR-002" not in result.waivers
    finally:
        os.unlink(path)


def test_coverage_goals_loaded_into_cov_items() -> None:
    """Coverage_Goals rows are loaded into cov_items keyed by COV ID."""
    cov_rows = [_simple_cov_row("COV-001", name="Reset coverage")]
    path = _make_workbook([], cov_rows)
    try:
        result = parse_vplan(path)
        assert "COV-001" in result.cov_items
    finally:
        os.unlink(path)


def test_cov_item_fields() -> None:
    """COV item entries contain all expected field keys."""
    cov_rows = [_simple_cov_row(
        "COV-002",
        name="TX FIFO full",
        description="Transmit FIFO reaches capacity",
        stimulus="Send maximum bytes",
        boundary="depth=16",
        linked="UART-FF-001, UART-FF-002",
        status="Open",
        cov_type="Functional",
        notes="Check overrun flag",
    )]
    path = _make_workbook([], cov_rows)
    try:
        result = parse_vplan(path)
        item = result.cov_items["COV-002"]
        assert item["name"] == "TX FIFO full"
        assert item["description"] == "Transmit FIFO reaches capacity"
        assert item["stimulus_strategy"] == "Send maximum bytes"
        assert item["boundary_values"] == "depth=16"
        assert item["status"] == "Open"
        assert item["coverage_type"] == "Functional"
        assert item["notes"] == "Check overrun flag"
    finally:
        os.unlink(path)


def test_linked_requirements_splits_on_comma() -> None:
    """Linked_Requirements with multiple IDs is split into a list."""
    cov_rows = [_simple_cov_row("COV-003", linked="UART-FF-001, UART-FF-002, UART-FF-003")]
    path = _make_workbook([], cov_rows)
    try:
        result = parse_vplan(path)
        assert result.cov_items["COV-003"]["linked_requirements"] == [
            "UART-FF-001", "UART-FF-002", "UART-FF-003"
        ]
    finally:
        os.unlink(path)


def test_req_ids_excludes_waived() -> None:
    """req_ids contains only non-waived requirement IDs."""
    rows = [
        _simple_vpr_row("UART-BR-001"),
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED", waiver_reason="OOS"),
        _simple_vpr_row("UART-BR-002"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert "UART-BR-001" in result.req_ids
        assert "UART-BR-002" in result.req_ids
        assert "UART-PAR-003" not in result.req_ids
    finally:
        os.unlink(path)


def test_req_schemes_derived_correctly() -> None:
    """req_schemes contains unique scheme prefixes from non-waived req_ids."""
    rows = [
        _simple_vpr_row("UART-BR-001"),
        _simple_vpr_row("UART-BR-002"),
        _simple_vpr_row("UART-PAR-001"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert "UART-BR" in result.req_schemes
        assert "UART-PAR" in result.req_schemes
        assert result.req_schemes.count("UART-BR") == 1
    finally:
        os.unlink(path)


def test_mode_full_when_non_waived_reqs_present() -> None:
    """mode == 'full' when at least one non-waived requirement exists."""
    rows = [
        _simple_vpr_row("UART-BR-001"),
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED", waiver_reason="OOS"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert result.mode == "full"
    finally:
        os.unlink(path)


def test_mode_campaign_when_all_waived() -> None:
    """mode == 'campaign' when all requirements are waived."""
    rows = [
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED", waiver_reason="OOS"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert result.mode == "campaign"
    finally:
        os.unlink(path)


def test_mode_campaign_when_empty() -> None:
    """mode == 'campaign' when no requirement rows are present."""
    path = _make_workbook([], [])
    try:
        result = parse_vplan(path)
        assert result.mode == "campaign"
    finally:
        os.unlink(path)


def test_req_parse_result_duck_type() -> None:
    """VplanParseResult duck-types ReqParseResult: requirements, waivers, mode."""
    rows = [
        _simple_vpr_row("UART-BR-001", statement="Stmt1", verif="simulation",
                         covered_by="COV-001"),
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED", waiver_reason="OOS"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        # requirements is a dict
        assert isinstance(result.requirements, dict)
        # each entry has the required keys
        entry = result.requirements["UART-BR-001"]
        assert "statement" in entry
        assert "verification" in entry
        assert isinstance(entry["verification"], list)
        assert "waived" in entry
        assert "waiver_reason" in entry
        # waivers is a list
        assert isinstance(result.waivers, list)
        # mode is a str
        assert isinstance(result.mode, str)
    finally:
        os.unlink(path)


def test_intent_parse_result_duck_type() -> None:
    """VplanParseResult duck-types IntentParseResult fields."""
    rows = [_simple_vpr_row("UART-BR-001")]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        # req_ids: list[str]
        assert isinstance(result.req_ids, list)
        # req_schemes: list[str]
        assert isinstance(result.req_schemes, list)
        # sections: dict (always empty)
        assert isinstance(result.sections, dict)
        assert result.sections == {}
        # inline_requirements: dict (always empty)
        assert isinstance(result.inline_requirements, dict)
        assert result.inline_requirements == {}
        # intent_waivers: list of dicts with item/reason/req_ids
        assert isinstance(result.intent_waivers, list)
    finally:
        os.unlink(path)


def test_intent_waivers_structure() -> None:
    """intent_waivers entries have item, reason, and req_ids keys."""
    rows = [
        _simple_vpr_row("UART-PAR-003", disposition="WAIVED",
                         waiver_reason="Not applicable"),
    ]
    path = _make_workbook(rows, [])
    try:
        result = parse_vplan(path)
        assert len(result.intent_waivers) == 1
        w = result.intent_waivers[0]
        assert w["item"] == "UART-PAR-003"
        assert w["reason"] == "Not applicable"
        assert w["req_ids"] == ["UART-PAR-003"]
    finally:
        os.unlink(path)


# ── Integration tests (BALU VPR fixture) ──────────────────────────────────────

_balu_missing = not os.path.exists(_BALU_VPLAN)
_balu_skip = pytest.mark.skipif(_balu_missing, reason="balu_vplan.xlsx fixture not present")


@_balu_skip
def test_balu_vplan_requirement_count() -> None:
    """BALU VPR fixture has exactly 141 data requirement rows."""
    result = parse_vplan(_BALU_VPLAN)
    assert len(result.requirements) == 141


@_balu_skip
def test_balu_vplan_waivers() -> None:
    """BALU VPR fixture has exactly 2 waived requirements: PAR-003 and VER-001."""
    result = parse_vplan(_BALU_VPLAN)
    assert len(result.waivers) == 2
    assert "UART-PAR-003" in result.waivers
    assert "UART-VER-001" in result.waivers


@_balu_skip
def test_balu_vplan_cov_item_count() -> None:
    """BALU VPR fixture has exactly 19 coverage items."""
    result = parse_vplan(_BALU_VPLAN)
    assert len(result.cov_items) == 19


@_balu_skip
def test_balu_vplan_intentional_gap() -> None:
    """UART-BR-004 has covered_by == '' and waived == False (intentional gap demo)."""
    result = parse_vplan(_BALU_VPLAN)
    entry = result.requirements.get("UART-BR-004")
    assert entry is not None
    assert entry["covered_by"] == ""
    assert entry["waived"] is False


@_balu_skip
def test_balu_vplan_mode_full() -> None:
    """BALU VPR fixture mode is 'full' (non-waived requirements present)."""
    result = parse_vplan(_BALU_VPLAN)
    assert result.mode == "full"
